"""Deterministic report workbook rendering (REQ-030/059/062)."""

from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

TOP_BLOCK_ROWS = (
    '统计周期与查询条件',
    '数据来源与统计口径',
    '质量摘要',
    '系统统计时钟生成时间与版本摘要',
    '签名短码与复核提示',
)
_TWO_DECIMAL_FIELDS = frozenset(
    {
        'usageQty',
        'peakQty',
        'flatQty',
        'valleyQty',
        'peakCost',
        'flatCost',
        'valleyCost',
        'totalCost',
        'coverageRatio',
        'oldValue',
        'newValue',
        'deltaValue',
        'deltaPct',
        'momPct',
        'effectiveRate',
        'savingValue',
        'price',
    }
)
_INTEGER_FIELDS = frozenset(
    {
        'count',
        'recordCount',
        'affectedObjectCount',
        'mergedCount',
    }
)


class ReportExcelService:
    """Render only the supplied canonical snapshot; never query business data."""

    @classmethod
    def build_workbook(
        cls,
        payload: dict[str, Any],
        full_signature: str,
    ) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        generated_at = datetime.fromisoformat(str(payload['generatedAt']))
        workbook.properties.created = generated_at
        workbook.properties.modified = generated_at
        basis = cls._basis(payload, full_signature)
        for section_name, section in payload['sections'].items():
            worksheet = workbook.create_sheet(section_name)
            cls._write_business_sheet(worksheet, section, basis)
        cls._write_basis_sheet(workbook.create_sheet('口径说明'), basis)
        output = BytesIO()
        archive = ZipFile(output, 'w', ZIP_DEFLATED, allowZip64=True)
        ExcelWriter(workbook, archive).save()
        return output.getvalue()

    @classmethod
    def deterministic_filename(cls, payload: dict[str, Any]) -> str:
        filters = payload.get('filters', {})
        parts = (
            payload.get('templateCode', 'REPORT'),
            payload.get('period', {}).get('label', 'period'),
            filters.get('zone', 'ALL'),
            filters.get('energyType', 'all'),
        )
        return '_'.join(str(part) for part in parts) + '.xlsx'

    @classmethod
    def _basis(
        cls,
        payload: dict[str, Any],
        full_signature: str,
    ) -> dict[str, str]:
        quality = payload.get('qualitySummary', {})
        versions = payload.get('versionSnapshots', {})
        return {
            'dataSources': cls._data_sources(payload),
            'period': cls._json(payload.get('period', {})),
            'queryConditions': cls._json(payload.get('filters', {})),
            'calculationNotes': cls._json(quality.get('calculationNotes', [])),
            'qualitySummary': cls._json(quality),
            'generatedAt': f"{payload.get('generatedAt')}（系统统计时钟）",
            'versionSnapshots': cls._json(versions),
            'fullSignature': full_signature,
            'verificationMethod': '同参数重导出比对签名',
        }

    @classmethod
    def _write_business_sheet(
        cls,
        worksheet: Any,
        section: Any,
        basis: dict[str, str],
    ) -> None:
        short_signature = basis['fullSignature'].split(':', 1)[-1][:12]
        top_values = (
            f"{basis['period']}；{basis['queryConditions']}",
            basis['dataSources'],
            basis['qualitySummary'],
            f"{basis['generatedAt']}；{basis['versionSnapshots']}",
            f'{short_signature}；{basis["verificationMethod"]}',
        )
        for row, (label, value) in enumerate(
            zip(TOP_BLOCK_ROWS, top_values, strict=True), start=1
        ):
            worksheet.cell(row, 1, label)
            worksheet.cell(row, 2, value)

        rows = cls._section_rows(section)
        columns = sorted({key for row in rows for key in row}) or ['value']
        for column, key in enumerate(columns, start=1):
            worksheet.cell(6, column, key)
        for row_number, item in enumerate(rows, start=7):
            for column, key in enumerate(columns, start=1):
                cell = worksheet.cell(
                    row_number,
                    column,
                    cls._cell_value(key, item.get(key)),
                )
                if key in _TWO_DECIMAL_FIELDS and item.get(key) is not None:
                    cell.number_format = '0.00'
                elif key in _INTEGER_FIELDS and item.get(key) is not None:
                    cell.number_format = '0'
        last_column = get_column_letter(len(columns))
        worksheet.auto_filter.ref = f'A6:{last_column}{max(6 + len(rows), 6)}'
        worksheet.freeze_panes = 'A7'

    @staticmethod
    def _write_basis_sheet(worksheet: Any, basis: dict[str, str]) -> None:
        for row, (key, value) in enumerate(basis.items(), start=1):
            worksheet.cell(row, 1, key)
            worksheet.cell(row, 2, value)

    @staticmethod
    def _section_rows(section: Any) -> list[dict[str, Any]]:
        if isinstance(section, dict) and isinstance(section.get('items'), list):
            rows = [item for item in section['items'] if isinstance(item, dict)]
            if rows:
                return rows
        if isinstance(section, dict):
            return [section]
        return [{'value': section}]

    @classmethod
    def _data_sources(cls, payload: dict[str, Any]) -> str:
        sections = payload.get('sections', {})
        sources: list[str] = []
        if 'usageSection' in sections:
            table = 'e_stat_day' if payload.get('templateCode') == 'ENERGY_DAILY' else 'e_stat_month'
            sources.append(table)
        if 'costSection' in sections:
            cost_versions = sorted(
                {
                    str(item.get('costVersion'))
                    for item in payload.get('versionSnapshots', {}).get('costVersions', [])
                    if isinstance(item, dict) and item.get('costVersion')
                }
            )
            if payload.get('templateCode') == 'COST_DIFF':
                sources.append(
                    'e_cost_recompute_record referenced costVersion: '
                    + ','.join(cost_versions)
                )
            else:
                sources.append(
                    f"e_cost_record current costVersion: {','.join(cost_versions)}"
                )
        if 'alertSection' in sections:
            sources.append('alert statistics（firstOccurredAt 落月）')
        if 'suggestionSection' in sections:
            sources.append('suggestions retrospective')
        if 'qualitySection' in sections:
            sources.append('statistics quality')
        return '；'.join(sources)

    @staticmethod
    def _cell_value(field: str, value: Any) -> Any:
        if value is not None and field in _TWO_DECIMAL_FIELDS:
            return float(value)
        if value is not None and field in _INTEGER_FIELDS:
            return int(value)
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return ReportExcelService._json(value)

    @staticmethod
    def _json(value: Any) -> str:
        return json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(',', ':'),
        )
