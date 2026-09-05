"""Act 5 canonical report assembly and preview contracts (REQ-030/059/061/062)."""

from __future__ import annotations

import importlib
import json
import os
import unittest
from copy import deepcopy
from dataclasses import replace
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from module_energy.controller.cost_controller import CostAccessGuard
from module_energy.entity.do.alert_event_do import EAlertEvent
from module_energy.entity.do.cost_recompute_record_do import ECostRecomputeRecord
from module_energy.entity.do.cost_record_do import ECostRecord
from module_energy.entity.do.report_archive_do import EReportArchive
from module_energy.entity.do.suggestion_do import ESuggestion
from module_energy.service.alert_service import AlertService
from module_energy.service.suggestion_service import SuggestionService


def _report_domain() -> object:
    try:
        return importlib.import_module('module_energy.domain.report_canonical')
    except ModuleNotFoundError as exc:
        raise AssertionError('report_canonical domain is required') from exc


def _report_vo() -> object:
    try:
        return importlib.import_module('module_energy.entity.vo.report_vo')
    except ModuleNotFoundError as exc:
        raise AssertionError('report VO is required') from exc


def _report_service() -> object:
    try:
        return importlib.import_module('module_energy.service.report_service').ReportService
    except ModuleNotFoundError as exc:
        raise AssertionError('ReportService is required') from exc


def _report_excel_service() -> object:
    try:
        return importlib.import_module(
            'module_energy.service.report_excel_service'
        ).ReportExcelService
    except ModuleNotFoundError as exc:
        raise AssertionError('ReportExcelService is required') from exc


class ReportDomainContractTest(unittest.TestCase):
    def test_signature_covers_sections_filters_versions_and_generated_at(self) -> None:
        domain = _report_domain()
        base = domain.CanonicalReportPayload(
            template_code='ENERGY_MONTHLY',
            template_version='v1',
            period={'start': '2026-06-01', 'end': '2026-07-01', 'state': 'complete'},
            filters={'zone': 'ALL', 'energyType': 'electricity'},
            generated_at=datetime(2026, 7, 12, 23, 59),
            sections={'costSection': {'items': [{'objectId': 1, 'totalCost': Decimal('1.20')}]}},
            quality_summary={'coverageRatio': Decimal('99.00')},
            version_snapshots={'costVersions': ['v1']},
        )
        signature = domain.build_report_signature(base)
        self.assertRegex(signature, r'^REPORT-SHA256-V1:[0-9a-f]{64}$')
        self.assertEqual(
            domain.canonical_report_bytes(base),
            domain.canonical_report_bytes(deepcopy(base)),
        )
        for field, changed in (
            ('sections', {'costSection': {'items': [{'objectId': 1, 'totalCost': Decimal('1.21')}]}}),
            ('filters', {'zone': 'B', 'energyType': 'electricity'}),
            ('generated_at', datetime(2026, 7, 12, 23, 59, 1)),
            ('version_snapshots', {'costVersions': ['v2']}),
        ):
            payload = replace(base, **{field: changed})
            self.assertNotEqual(signature, domain.build_report_signature(payload), field)

    def test_request_is_strict_camel_case_and_rejects_invalid_periods(self) -> None:
        vo = _report_vo()
        valid = vo.ReportPreviewRequest.model_validate(
            {
                'templateCode': 'ENERGY_MONTHLY',
                'period': '2026-06',
                'filters': {'zone': 'ALL', 'energyType': 'electricity'},
            }
        )
        self.assertEqual('ENERGY_MONTHLY', valid.template_code)
        for invalid in (
            {'template_code': 'ENERGY_MONTHLY'},
            {'templateCode': 'UNKNOWN'},
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-6'},
            {'templateCode': 'ENERGY_DAILY', 'period': '2026-07'},
            {'templateCode': 'ENERGY_DAILY', 'period': '2026-99-99'},
            {'templateCode': 'ENERGY_DAILY', 'period': '2026-02-30'},
            {'templateCode': 'EQUIPMENT_PROFILE', 'period': '2026-07'},
            {
                'templateCode': 'ENERGY_MONTHLY',
                'filters': {'equipmentId': 1},
            },
            {'templateCode': 'ENERGY_DAILY', 'unexpected': True},
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                vo.ReportPreviewRequest.model_validate(invalid)

    def test_excel_has_five_row_header_filter_and_independent_basis_sheet(self) -> None:
        domain = _report_domain()
        payload = domain.CanonicalReportPayload(
            template_code='ENERGY_MONTHLY',
            template_version='v1',
            period={
                'label': '2026-06',
                'start': '2026-06-01',
                'end': '2026-07-01',
                'state': 'complete',
            },
            filters={'zone': 'ALL', 'energyType': 'electricity'},
            generated_at=datetime(2026, 7, 12, 23, 59),
            sections={
                'usageSection': {
                    'items': [
                        {
                            'objectId': 1,
                            'usageQty': Decimal('10.20'),
                            'coverageRatio': Decimal('99.00'),
                        }
                    ]
                },
                'costSection': {
                    'items': [
                        {
                            'objectId': 1,
                            'totalCost': Decimal('8.10'),
                            'currentCostVersion': 'v1',
                        }
                    ]
                },
            },
            quality_summary={
                'coverageRatio': Decimal('99.00'),
                'calculationNotes': [],
            },
            version_snapshots={
                'costVersions': [{'costVersion': 'v1'}],
                'tariffVersions': [{'versionNo': 1, 'touPeriod': 'peak'}],
                'allocRuleVersions': [{'versionNo': 1}],
            },
        )
        signature = domain.build_report_signature(payload)

        workbook_bytes = _report_excel_service().build_workbook(
            payload.to_dict(), signature
        )
        workbook = load_workbook(BytesIO(workbook_bytes))

        self.assertEqual(
            {'usageSection', 'costSection', '口径说明'}, set(workbook.sheetnames)
        )
        self.assertEqual(datetime(2026, 7, 12, 23, 59), workbook.properties.created)
        self.assertEqual(datetime(2026, 7, 12, 23, 59), workbook.properties.modified)
        expected_labels = [
            '统计周期与查询条件',
            '数据来源与统计口径',
            '质量摘要',
            '系统统计时钟生成时间与版本摘要',
            '签名短码与复核提示',
        ]
        for sheet_name in ('usageSection', 'costSection'):
            sheet = workbook[sheet_name]
            self.assertEqual(expected_labels, [sheet.cell(row, 1).value for row in range(1, 6)])
            self.assertEqual('A7', sheet.freeze_panes)
            self.assertIsNotNone(sheet.auto_filter.ref)
            self.assertEqual([], list(sheet.merged_cells.ranges))
            self.assertTrue(any(sheet.cell(6, column).value for column in range(1, sheet.max_column + 1)))

        usage_sheet = workbook['usageSection']
        usage_headers = {
            usage_sheet.cell(6, column).value: column
            for column in range(1, usage_sheet.max_column + 1)
        }
        for field, expected in (('usageQty', 10.2), ('coverageRatio', 99.0)):
            cell = usage_sheet.cell(7, usage_headers[field])
            self.assertEqual('n', cell.data_type)
            self.assertEqual('0.00', cell.number_format)
            self.assertEqual(expected, cell.value)
        object_id_cell = usage_sheet.cell(7, usage_headers['objectId'])
        self.assertEqual('n', object_id_cell.data_type)
        self.assertEqual('General', object_id_cell.number_format)
        self.assertEqual(1, object_id_cell.value)

        cost_sheet = workbook['costSection']
        cost_headers = {
            cost_sheet.cell(6, column).value: column
            for column in range(1, cost_sheet.max_column + 1)
        }
        total_cost_cell = cost_sheet.cell(7, cost_headers['totalCost'])
        self.assertEqual('n', total_cost_cell.data_type)
        self.assertEqual('0.00', total_cost_cell.number_format)
        self.assertEqual(8.1, total_cost_cell.value)
        self.assertEqual(
            'v1', cost_sheet.cell(7, cost_headers['currentCostVersion']).value
        )

        basis = workbook['口径说明']
        basis_values = {
            basis.cell(row, 1).value: basis.cell(row, 2).value
            for row in range(1, basis.max_row + 1)
        }
        self.assertEqual(
            {
                'dataSources',
                'period',
                'queryConditions',
                'calculationNotes',
                'qualitySummary',
                'generatedAt',
                'versionSnapshots',
                'fullSignature',
                'verificationMethod',
            },
            set(basis_values),
        )
        self.assertIn('current costVersion: v1', basis_values['dataSources'])
        self.assertIn('系统统计时钟', basis_values['generatedAt'])
        self.assertEqual(signature, basis_values['fullSignature'])
        self.assertIn('同参数重导出比对签名', basis_values['verificationMethod'])
        self.assertIn(signature.split(':', 1)[1][:12], workbook['costSection']['B5'].value)

        controller = importlib.import_module('module_energy.controller.report_controller')
        response = controller._excel_response(
            {
                'content': workbook_bytes,
                'filename': 'ENERGY_MONTHLY_2026-06_ALL_electricity.xlsx',
                'signature': signature,
            }
        )
        self.assertIn(
            'ENERGY_MONTHLY_2026-06_ALL_electricity.xlsx',
            response.headers['Content-Disposition'],
        )
        self.assertEqual(signature, response.headers['X-Report-Signature'])
        self.assertEqual(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response.media_type,
        )

    def test_cost_diff_basis_labels_old_and_new_versions_as_referenced(self) -> None:
        payload = {
            'templateCode': 'COST_DIFF',
            'templateVersion': 'v1',
            'period': {
                'label': '2026-06',
                'start': '2026-06-01',
                'end': '2026-07-01',
                'state': 'complete',
            },
            'filters': {'zone': 'ALL', 'energyType': 'electricity'},
            'generatedAt': '2026-07-12T23:59:00',
            'sections': {'costSection': {'items': []}},
            'qualitySummary': {'calculationNotes': []},
            'versionSnapshots': {
                'costVersions': [
                    {'costVersion': 'v1'},
                    {'costVersion': 'v2'},
                ]
            },
        }

        workbook = load_workbook(
            BytesIO(
                _report_excel_service().build_workbook(
                    payload,
                    'REPORT-SHA256-V1:' + '0' * 64,
                )
            )
        )
        basis = workbook['口径说明']
        basis_values = {
            basis.cell(row, 1).value: basis.cell(row, 2).value
            for row in range(1, basis.max_row + 1)
        }

        self.assertIn(
            'e_cost_recompute_record referenced costVersion: v1,v2',
            basis_values['dataSources'],
        )
        self.assertNotIn('current costVersion', basis_values['dataSources'])


class ReportControllerContractTest(unittest.TestCase):
    def test_router_exposes_task7_endpoints_with_cost_guard(self) -> None:
        try:
            controller = importlib.import_module('module_energy.controller.report_controller')
        except ModuleNotFoundError as exc:
            self.fail(f'report controller is required: {exc}')
        routes = {
            (route.path, frozenset(route.methods or set())): route
            for route in controller.report_controller.routes
        }
        self.assertEqual(
            {
                ('/reports/templates', frozenset({'GET'})),
                ('/reports/preview', frozenset({'POST'})),
                ('/reports/export', frozenset({'POST'})),
                ('/reports/archives', frozenset({'POST'})),
                ('/reports/archives', frozenset({'GET'})),
                ('/reports/archives/{archive_id}', frozenset({'GET'})),
                ('/reports/archives/{archive_id}/export', frozenset({'GET'})),
            },
            set(routes),
        )
        for route in routes.values():
            guards = [
                dependency.dependency
                for dependency in route.dependencies
                if isinstance(dependency.dependency, CostAccessGuard)
            ]
            self.assertEqual(1, len(guards))
            self.assertEqual(frozenset({'finance', 'energy_mgr'}), guards[0].allowed_roles)
            self.assertFalse(guards[0].allow_admin)


class ReportIntegrationTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        test_db = os.environ.get('ACT5_TEST_DB', '')
        self.assertTrue(test_db, 'ACT5_TEST_DB must be set; Task 6 may not skip')
        self.assertEqual(test_db, os.environ.get('ACT4_TEST_DB'))
        self.assertRegex(test_db, r'^codex_[a-z0-9_]*_test$')
        self.engine = create_async_engine(
            f'mysql+asyncmy://demo:bdemo_dev@127.0.0.1:3306/{test_db}'
        )
        self.session_factory = async_sessionmaker(self.engine, expire_on_commit=False)
        self.db = self.session_factory()
        counts = (
            int((await self.db.execute(select(func.count()).select_from(ECostRecord))).scalar_one()),
            int((await self.db.execute(select(func.count()).select_from(EAlertEvent))).scalar_one()),
            int((await self.db.execute(select(func.count()).select_from(ESuggestion))).scalar_one()),
        )
        self.assertEqual((75, 20, 8), counts)
        self.before_counts = counts
        self.created_archive_ids: list[int] = []
        self.assertEqual(
            0,
            int(
                (
                    await self.db.execute(select(func.count()).select_from(EReportArchive))
                ).scalar_one()
            ),
            'Task 7 requires a fresh archive table',
        )

    async def asyncTearDown(self) -> None:
        await self.db.rollback()
        await self.db.close()
        if self.created_archive_ids:
            async with self.session_factory() as cleanup_db:
                await cleanup_db.execute(
                    delete(EReportArchive).where(
                        EReportArchive.archive_id.in_(self.created_archive_ids)
                    )
                )
                await cleanup_db.commit()
        await self.engine.dispose()

    async def _counts(self) -> tuple[int, int, int]:
        return (
            int((await self.db.execute(select(func.count()).select_from(ECostRecord))).scalar_one()),
            int((await self.db.execute(select(func.count()).select_from(EAlertEvent))).scalar_one()),
            int((await self.db.execute(select(func.count()).select_from(ESuggestion))).scalar_one()),
        )

    async def test_monthly_default_is_deterministic_and_reuses_shared_services(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY'}
        )
        domain = _report_domain()
        with (
            patch.object(
                AlertService,
                'get_statistics',
                wraps=AlertService.get_statistics,
            ) as alert_spy,
            patch.object(
                SuggestionService,
                'get_retrospective',
                wraps=SuggestionService.get_retrospective,
            ) as suggestion_spy,
        ):
            first = await service.assemble(self.db, request)
            second = await service.assemble(self.db, request)

        self.assertEqual(first.to_dict(), second.to_dict())
        self.assertEqual(
            domain.canonical_report_bytes(first),
            domain.canonical_report_bytes(second),
        )
        self.assertEqual(
            domain.build_report_signature(first),
            domain.build_report_signature(second),
        )
        self.assertEqual('2026-06', first.period['label'])
        self.assertEqual('2026-06-01', first.period['start'])
        self.assertEqual('2026-07-01', first.period['end'])
        self.assertEqual('complete', first.period['state'])
        self.assertEqual(datetime(2026, 7, 12).date(), first.generated_at.date())
        self.assertEqual(
            {
                'usageSection',
                'costSection',
                'alertSection',
                'suggestionSection',
                'qualitySection',
            },
            set(first.sections),
        )
        alert_kwargs = alert_spy.call_args_list[0].kwargs
        self.assertEqual(datetime(2026, 6, 1), alert_kwargs['period_start'])
        self.assertEqual(datetime(2026, 7, 1), alert_kwargs['period_end'])
        suggestion_kwargs = suggestion_spy.call_args_list[0].kwargs
        self.assertEqual('2026-06', suggestion_kwargs['month'])
        self.assertEqual('ALL', suggestion_kwargs['zone'])
        self.assertIn('costVersions', first.version_snapshots)
        self.assertIn('tariffVersions', first.version_snapshots)
        self.assertIn('allocRuleVersions', first.version_snapshots)
        self.assertIn('templateVersions', first.version_snapshots)
        self.assertEqual(self.before_counts, await self._counts())

    async def test_daily_uses_day_window_and_prunes_cost_and_template_versions(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_DAILY'}
        )
        with (
            patch.object(
                AlertService,
                'get_statistics',
                wraps=AlertService.get_statistics,
            ) as alert_spy,
            patch.object(
                SuggestionService,
                'get_retrospective',
                wraps=SuggestionService.get_retrospective,
            ) as suggestion_spy,
        ):
            payload = await service.assemble(self.db, request)
        self.assertEqual('2026-07-12', payload.period['label'])
        self.assertEqual(
            {'usageSection', 'alertSection', 'qualitySection'},
            set(payload.sections),
        )
        kwargs = alert_spy.call_args.kwargs
        self.assertEqual(datetime(2026, 7, 12), kwargs['period_start'])
        self.assertEqual(datetime(2026, 7, 13), kwargs['period_end'])
        suggestion_spy.assert_not_called()
        forbidden = {'costVersions', 'tariffVersions', 'allocRuleVersions', 'templateVersions'}
        self.assertTrue(forbidden.isdisjoint(payload.version_snapshots))
        self.assertEqual(self.before_counts, await self._counts())

    async def test_in_progress_month_and_special_template_section_pruning(self) -> None:
        service = _report_service()
        monthly = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-07'}
        )
        in_progress = await service.assemble(self.db, monthly)
        self.assertEqual('inProgress', in_progress.period['state'])
        self.assertEqual('2026-07-12', in_progress.period['asOf'])

        profile = _report_vo().ReportPreviewRequest.model_validate(
            {
                'templateCode': 'EQUIPMENT_PROFILE',
                'period': '2026-07',
                'filters': {'equipmentId': 1},
            }
        )
        profile_payload = await service.assemble(self.db, profile)
        self.assertNotIn('costSection', profile_payload.sections)
        self.assertNotIn('suggestionSection', profile_payload.sections)
        forbidden = {'costVersions', 'tariffVersions', 'allocRuleVersions', 'templateVersions'}
        self.assertTrue(forbidden.isdisjoint(profile_payload.version_snapshots))

        for filters in ({'equipmentId': 999999}, {'zone': 'B', 'equipmentId': 1}):
            invalid_profile = _report_vo().ReportPreviewRequest.model_validate(
                {
                    'templateCode': 'EQUIPMENT_PROFILE',
                    'period': '2026-07',
                    'filters': filters,
                }
            )
            with self.subTest(filters=filters), self.assertRaises(HTTPException) as error:
                await service.assemble(self.db, invalid_profile)
            self.assertEqual(422, error.exception.status_code)

    async def test_versions_use_exact_participants_and_business_enum_order(self) -> None:
        service = _report_service()
        july = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-07'}
        )

        payload = await service.assemble(self.db, july)

        self.assertEqual(
            [5, 6, 7, 8],
            [item['suggestionId'] for item in payload.version_snapshots['templateVersions']],
        )

        june = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-06'}
        )
        june_payload = await service.assemble(self.db, june)
        self.assertEqual(
            ['peak', 'flat', 'valley'],
            [item['touPeriod'] for item in june_payload.version_snapshots['tariffVersions']],
        )

    async def test_may_discloses_partial_data_window(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-05'}
        )

        payload = await service.assemble(self.db, request)

        self.assertEqual('partial', payload.period['state'])
        self.assertEqual('2026-05-04', payload.period['dataStart'])
        self.assertEqual('2026-05-31', payload.period['dataEnd'])
        expected_note = '2026-05 为 05-04 起的部分月，环比分母为已覆盖周期金额'
        self.assertIn(expected_note, payload.quality_summary['notes'])
        self.assertIn(expected_note, payload.sections['qualitySection']['calculationNotes'])

    async def test_cost_diff_is_honestly_empty_without_recompute_evidence(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'COST_DIFF', 'period': '2026-06'}
        )

        payload = await service.assemble(self.db, request)

        self.assertEqual({'items': []}, payload.sections['costSection'])
        self.assertEqual({}, payload.version_snapshots)

        current = (
            await self.db.execute(
                select(ECostRecord).where(
                    ECostRecord.stat_month == '2026-06',
                    ECostRecord.energy_type_code == 'electricity',
                    ECostRecord.object_type == 'area',
                    ECostRecord.object_id == 1,
                    ECostRecord.is_current.is_(True),
                )
            )
        ).scalar_one()
        record = ECostRecomputeRecord(
            period_key='2026-06:ALL:electricity',
            stat_month='2026-06',
            energy_type_code='electricity',
            scope='ALL',
            old_cost_version=1,
            new_cost_version=2,
            trigger_reason='测试单价版本生效',
            trigger_type='tariffVersionActivated',
            triggered_by='finance_user',
            triggered_at=datetime(2026, 7, 12, 10, 0),
            tariff_snapshot_json=current.tariff_snapshot_json,
            alloc_rule_snapshot_json=current.alloc_rule_snapshot_json,
            diff_summary_json=json.dumps(
                [
                    {
                        'objectType': 'system',
                        'objectId': None,
                        'objectCode': 'SYSTEM',
                        'metric': 'totalCost',
                        'oldValue': 30.005,
                        'newValue': 33.004,
                        'deltaValue': 2.999,
                        'deltaPct': None,
                        'oldSignature': 'old-system',
                        'newSignature': 'new-system',
                    },
                    {
                        'objectType': 'area',
                        'objectId': 2,
                        'objectCode': 'AREA-B',
                        'metric': 'totalCost',
                        'oldValue': 10.1,
                        'newValue': 12.0,
                        'deltaValue': 1.9,
                        'deltaPct': 20.126,
                        'oldSignature': 'old-area-b',
                        'newSignature': 'new-area-b',
                    },
                    {
                        'objectType': 'area',
                        'objectId': 1,
                        'objectCode': 'AREA-A',
                        'metric': 'totalCost',
                        'oldValue': 20.0,
                        'newValue': 21.0,
                        'deltaValue': 1.0,
                        'deltaPct': 5.0,
                        'oldSignature': 'old-area-a',
                        'newSignature': 'new-area-a',
                    },
                    {
                        'objectType': 'equipment',
                        'objectId': 8,
                        'objectCode': 'PN-B1',
                        'metric': 'totalCost',
                        'oldValue': 5.555,
                        'newValue': 6.0,
                        'deltaValue': 0.445,
                        'deltaPct': 8.005,
                        'oldSignature': 'old-pn-b1',
                        'newSignature': 'new-pn-b1',
                    },
                ],
                ensure_ascii=False,
            ),
            review_status='approved',
            reviewed_by='finance_user',
            reviewed_at=datetime(2026, 7, 12, 11, 0),
            review_remark='测试复核通过',
        )
        self.db.add(record)
        await self.db.flush()

        with_evidence = await service.assemble(self.db, request)

        self.assertEqual(1, len(with_evidence.sections['costSection']['items']))
        item = with_evidence.sections['costSection']['items'][0]
        self.assertEqual(int(record.recompute_id), item['recomputeId'])
        self.assertEqual(
            {
                'recomputeId',
                'statMonth',
                'energyType',
                'scope',
                'oldCostVersion',
                'newCostVersion',
                'triggerReason',
                'triggerType',
                'triggeredBy',
                'triggeredAt',
                'reviewStatus',
                'reviewedBy',
                'reviewedAt',
                'reviewRemark',
                'diffSummary',
            },
            set(item),
        )
        self.assertEqual('v1', item['oldCostVersion'])
        self.assertEqual('v2', item['newCostVersion'])
        self.assertEqual('approved', item['reviewStatus'])
        self.assertEqual(
            [('area', 1), ('area', 2), ('equipment', 8), ('system', None)],
            [(part['objectType'], part['objectId']) for part in item['diffSummary']],
        )
        by_object = {
            (part['objectType'], part['objectId']): part
            for part in item['diffSummary']
        }
        self.assertEqual(Decimal('10.10'), by_object[('area', 2)]['oldValue'])
        self.assertEqual(Decimal('20.13'), by_object[('area', 2)]['deltaPct'])
        self.assertEqual(Decimal('5.56'), by_object[('equipment', 8)]['oldValue'])
        self.assertEqual(Decimal('8.01'), by_object[('equipment', 8)]['deltaPct'])
        self.assertEqual(Decimal('30.01'), by_object[('system', None)]['oldValue'])
        self.assertIsNone(by_object[('system', None)]['deltaPct'])
        self.assertTrue(
            all(
                set(part)
                == {
                    'objectType',
                    'objectId',
                    'objectCode',
                    'metric',
                    'oldValue',
                    'newValue',
                    'deltaValue',
                    'deltaPct',
                    'oldSignature',
                    'newSignature',
                }
                for part in item['diffSummary']
            )
        )
        self.assertTrue(
            {'objectId', 'objectCode', 'totalCost', 'currentCostVersion'}.isdisjoint(item)
        )
        self.assertEqual(
            {'costVersions', 'tariffVersions', 'allocRuleVersions'},
            set(with_evidence.version_snapshots),
        )
        self.assertEqual(
            ['v1', 'v2'],
            [part['costVersion'] for part in with_evidence.version_snapshots['costVersions']],
        )
        self.assertEqual(
            ['peak', 'flat', 'valley'],
            [part['touPeriod'] for part in with_evidence.version_snapshots['tariffVersions']],
        )
        self.assertEqual(
            [json.loads(current.alloc_rule_snapshot_json)],
            with_evidence.version_snapshots['allocRuleVersions'],
        )
        serialized = with_evidence.to_dict()
        serialized_item = serialized['sections']['costSection']['items'][0]
        serialized_by_object = {
            (part['objectType'], part['objectId']): part
            for part in serialized_item['diffSummary']
        }
        self.assertEqual('10.10', serialized_by_object[('area', 2)]['oldValue'])
        self.assertEqual('20.13', serialized_by_object[('area', 2)]['deltaPct'])
        self.assertEqual('5.56', serialized_by_object[('equipment', 8)]['oldValue'])
        self.assertEqual('8.01', serialized_by_object[('equipment', 8)]['deltaPct'])
        self.assertEqual('30.01', serialized_by_object[('system', None)]['oldValue'])
        self.assertIsNone(serialized_by_object[('system', None)]['deltaPct'])
        self.assertEqual(
            serialized,
            json.loads(_report_domain().canonical_report_bytes(with_evidence)),
        )

        area_a_only = ECostRecomputeRecord(
            period_key='2026-06:AREA-A:electricity',
            stat_month='2026-06',
            energy_type_code='electricity',
            scope='AREA-A',
            old_cost_version=1,
            new_cost_version=2,
            trigger_reason='仅 A 区差异',
            trigger_type='manual',
            triggered_by='finance_user',
            triggered_at=datetime(2026, 7, 12, 12, 0),
            tariff_snapshot_json=current.tariff_snapshot_json,
            alloc_rule_snapshot_json=current.alloc_rule_snapshot_json,
            diff_summary_json=json.dumps([serialized_item['diffSummary'][0]], ensure_ascii=False),
            review_status='pending',
        )
        self.db.add(area_a_only)
        await self.db.flush()
        zone_b_request = _report_vo().ReportPreviewRequest.model_validate(
            {
                'templateCode': 'COST_DIFF',
                'period': '2026-06',
                'filters': {'zone': 'B'},
            }
        )

        zone_b = await service.assemble(self.db, zone_b_request)

        self.assertEqual(1, len(zone_b.sections['costSection']['items']))
        self.assertEqual(
            [('area', 2), ('equipment', 8)],
            [
                (part['objectType'], part['objectId'])
                for part in zone_b.sections['costSection']['items'][0]['diffSummary']
            ],
        )

    async def test_empty_usage_has_deterministic_quality_note(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {
                'templateCode': 'EQUIPMENT_PROFILE',
                'period': '2026-07',
                'filters': {'equipmentId': 1, 'energyType': 'compressed_air'},
            }
        )

        payload = await service.assemble(self.db, request)

        self.assertEqual(0, payload.quality_summary['recordCount'])
        self.assertEqual(['所选周期无可用统计数据'], payload.quality_summary['notes'])

    async def test_export_reuses_assemble_and_keeps_daily_versions_truthful(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_DAILY'}
        )
        with patch.object(service, 'assemble', wraps=service.assemble) as assemble_spy:
            first = await service.export(self.db, request)
            second = await service.export(self.db, request)

        self.assertEqual(2, assemble_spy.await_count)
        self.assertEqual(first['signature'], second['signature'])
        self.assertEqual(first['filename'], second['filename'])
        self.assertEqual(
            'ENERGY_DAILY_2026-07-12_ALL_electricity.xlsx', first['filename']
        )
        self.assertRegex(first['signature'], r'^REPORT-SHA256-V1:[0-9a-f]{64}$')
        workbook = load_workbook(BytesIO(first['content']))
        basis = workbook['口径说明']
        basis_values = {
            basis.cell(row, 1).value: basis.cell(row, 2).value
            for row in range(1, basis.max_row + 1)
        }
        self.assertNotIn('costVersion', basis_values['dataSources'])
        self.assertNotIn('costVersions', basis_values['versionSnapshots'])
        self.assertNotIn('tariffVersions', basis_values['versionSnapshots'])
        self.assertNotIn('allocRuleVersions', basis_values['versionSnapshots'])
        self.assertEqual(
            0,
            int(
                (
                    await self.db.execute(select(func.count()).select_from(EReportArchive))
                ).scalar_one()
            ),
        )

    async def test_pending_cost_blocks_formal_output_without_archive_write(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-06'}
        )
        current = (
            await self.db.execute(
                select(ECostRecord).where(
                    ECostRecord.stat_month == '2026-06',
                    ECostRecord.energy_type_code == 'electricity',
                    ECostRecord.object_type == 'area',
                    ECostRecord.object_id == 1,
                    ECostRecord.is_current.is_(True),
                )
            )
        ).scalar_one()
        for status in ('draft', 'pendingReview', 'pendingRecompute'):
            current.status = status
            await self.db.flush()

            preview = await service.preview(self.db, request)

            self.assertEqual(
                status,
                preview['sections']['costSection']['reviewState'],
            )
            with (
                patch.object(
                    _report_excel_service(),
                    'build_workbook',
                    wraps=_report_excel_service().build_workbook,
                ) as workbook_spy,
                self.assertRaises(HTTPException) as export_error,
            ):
                await service.export(self.db, request)
            self.assertEqual(409, export_error.exception.status_code)
            workbook_spy.assert_not_called()
            with self.assertRaises(HTTPException) as archive_error:
                await service.create_archive(
                    self.db,
                    request,
                    archived_by='finance_user',
                )
            self.assertEqual(409, archive_error.exception.status_code)

        current.status = 'frozen'
        await self.db.flush()
        frozen_file = await service.export(self.db, request)
        self.assertRegex(frozen_file['signature'], r'^REPORT-SHA256-V1:[0-9a-f]{64}$')
        current.status = 'reviewed'
        pending_diff = ECostRecomputeRecord(
            period_key='2026-06:ALL:electricity',
            stat_month='2026-06',
            energy_type_code='electricity',
            scope='ALL',
            old_cost_version=1,
            new_cost_version=2,
            trigger_reason='待复核差异专项',
            trigger_type='tariffVersionActivated',
            triggered_by='finance_user',
            triggered_at=datetime(2026, 7, 12, 12, 0),
            tariff_snapshot_json=current.tariff_snapshot_json,
            alloc_rule_snapshot_json=current.alloc_rule_snapshot_json,
            diff_summary_json=json.dumps(
                [
                    {
                        'objectType': 'area',
                        'objectId': 1,
                        'objectCode': 'AREA-A',
                        'metric': 'totalCost',
                        'oldValue': 10.0,
                        'newValue': 11.0,
                        'deltaValue': 1.0,
                        'deltaPct': 10.0,
                        'oldSignature': 'old',
                        'newSignature': 'new',
                    }
                ],
                ensure_ascii=False,
            ),
            review_status='pending',
        )
        self.db.add(pending_diff)
        await self.db.flush()
        diff_request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'COST_DIFF', 'period': '2026-06'}
        )
        diff_preview = await service.preview(self.db, diff_request)
        self.assertEqual(
            'pending',
            diff_preview['sections']['costSection']['items'][0]['reviewStatus'],
        )
        with (
            patch.object(
                _report_excel_service(),
                'build_workbook',
                wraps=_report_excel_service().build_workbook,
            ) as diff_workbook_spy,
            self.assertRaises(HTTPException) as diff_export_error,
        ):
            await service.export(self.db, diff_request)
        self.assertEqual(409, diff_export_error.exception.status_code)
        diff_workbook_spy.assert_not_called()
        with self.assertRaises(HTTPException) as diff_archive_error:
            await service.create_archive(
                self.db,
                diff_request,
                archived_by='finance_user',
            )
        self.assertEqual(409, diff_archive_error.exception.status_code)
        for reviewed_status in ('approved', 'rejected'):
            pending_diff.review_status = reviewed_status
            await self.db.flush()
            reviewed_file = await service.export(self.db, diff_request)
            self.assertRegex(
                reviewed_file['signature'], r'^REPORT-SHA256-V1:[0-9a-f]{64}$'
            )
        self.assertEqual(
            0,
            int(
                (
                    await self.db.execute(select(func.count()).select_from(EReportArchive))
                ).scalar_one()
            ),
        )

    async def test_archive_detail_and_export_read_only_frozen_payload(self) -> None:
        service = _report_service()
        monthly_request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-06'}
        )
        daily_request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_DAILY'}
        )
        initial_file = await service.export(self.db, monthly_request)
        monthly_archive = await service.create_archive(
            self.db,
            monthly_request,
            archived_by='finance_user',
        )
        self.created_archive_ids.append(monthly_archive['archiveId'])
        self.assertEqual(
            {
                'archiveId',
                'templateCode',
                'templateVersion',
                'periodStart',
                'periodEnd',
                'filtersSnapshot',
                'payloadSnapshot',
                'versionSnapshots',
                'fullSignature',
                'generatedAt',
                'archivedBy',
                'archivedAt',
            },
            set(monthly_archive),
        )
        self.assertEqual(initial_file['signature'], monthly_archive['fullSignature'])
        self.assertEqual(
            monthly_archive['generatedAt'], monthly_archive['archivedAt']
        )
        async with self.session_factory() as persisted_db:
            persisted = await service.get_archive(
                persisted_db,
                monthly_archive['archiveId'],
            )
        self.assertEqual(
            monthly_archive['fullSignature'], persisted['fullSignature']
        )
        daily_archive = await service.create_archive(
            self.db,
            daily_request,
            archived_by='energy_mgr_user',
        )
        self.created_archive_ids.append(daily_archive['archiveId'])
        archives = await service.list_archives(self.db)
        self.assertEqual(
            [daily_archive['archiveId'], monthly_archive['archiveId']],
            [item['archiveId'] for item in archives['items']],
        )

        current = (
            await self.db.execute(
                select(ECostRecord).where(
                    ECostRecord.stat_month == '2026-06',
                    ECostRecord.energy_type_code == 'electricity',
                    ECostRecord.object_type == 'area',
                    ECostRecord.object_id == 1,
                    ECostRecord.is_current.is_(True),
                )
            )
        ).scalar_one()
        current.total_cost = Decimal(current.total_cost) + Decimal('999.99')
        await self.db.flush()

        with patch.object(service, 'assemble', side_effect=AssertionError('must not replay')):
            detail = await service.get_archive(self.db, monthly_archive['archiveId'])
            frozen_file = await service.export_archive(
                self.db, monthly_archive['archiveId']
            )

        self.assertEqual(monthly_archive['fullSignature'], detail['fullSignature'])
        self.assertEqual(
            monthly_archive['payloadSnapshot'], detail['payloadSnapshot']
        )
        self.assertEqual(initial_file['signature'], frozen_file['signature'])
        self.assertEqual(initial_file['filename'], frozen_file['filename'])
        workbook = load_workbook(BytesIO(frozen_file['content']))
        cost_sheet = workbook['costSection']
        headers = {
            cost_sheet.cell(6, column).value: column
            for column in range(1, cost_sheet.max_column + 1)
        }
        self.assertEqual(
            float(
                monthly_archive['payloadSnapshot']['sections']['costSection']['items'][
                    0
                ]['totalCost']
            ),
            cost_sheet.cell(7, headers['totalCost']).value,
        )
        self.assertEqual(
            '0.00', cost_sheet.cell(7, headers['totalCost']).number_format
        )
        with self.assertRaises(HTTPException) as not_found:
            await service.get_archive(self.db, 999999)
        self.assertEqual(404, not_found.exception.status_code)

    async def test_archive_returns_with_production_expire_on_commit_session(self) -> None:
        service = _report_service()
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_DAILY'}
        )
        production_factory = async_sessionmaker(self.engine)
        async with production_factory() as production_db:
            self.assertTrue(production_db.sync_session.expire_on_commit)
            archive = await service.create_archive(
                production_db,
                request,
                archived_by='finance_user',
            )
        self.created_archive_ids.append(archive['archiveId'])

        async with production_factory() as reader_db:
            persisted = await service.get_archive(reader_db, archive['archiveId'])

        self.assertEqual(archive['fullSignature'], persisted['fullSignature'])
        self.assertEqual(archive['payloadSnapshot'], persisted['payloadSnapshot'])

    async def test_templates_and_preview_response_are_read_only(self) -> None:
        service = _report_service()
        templates = await service.list_templates(self.db)
        self.assertEqual(
            {
                'ENERGY_DAILY',
                'ENERGY_MONTHLY',
                'EQUIPMENT_PROFILE',
                'SUGGESTION_RETROSPECTIVE',
                'COST_DIFF',
            },
            {item['templateCode'] for item in templates['items']},
        )
        self.assertEqual(
            {
                'enabled': False,
                'priority': 'P2',
                'label': '报表订阅（规划中）',
            },
            templates['subscription'],
        )
        request = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY'}
        )
        preview = await service.preview(self.db, request)
        self.assertEqual({'reportMeta', 'sections', 'signature'}, set(preview))
        self.assertRegex(preview['signature'], r'^REPORT-SHA256-V1:[0-9a-f]{64}$')
        self.assertEqual(self.before_counts, await self._counts())

        future = _report_vo().ReportPreviewRequest.model_validate(
            {'templateCode': 'ENERGY_MONTHLY', 'period': '2026-08'}
        )
        with self.assertRaises(HTTPException) as invalid_period:
            await service.assemble(self.db, future)
        self.assertEqual(422, invalid_period.exception.status_code)
        self.assertEqual(self.before_counts, await self._counts())


if __name__ == '__main__':
    unittest.main()
