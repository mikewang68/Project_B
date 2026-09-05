"""Act 5 destructive end-to-end workflow and reset determinism gate.

REQ-030/051~056/059/062/073/074.  This suite owns a disposable MySQL
database named by identical ``ACT4_TEST_DB``/``ACT5_TEST_DB`` variables.  It
never touches ``b_demo`` and restores the authoritative reset/bootstrap
fixture even when an intermediate assertion fails (P-17/P-21).
"""

from __future__ import annotations

import asyncio
import inspect
import json
import os
import subprocess
import unittest
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import module_energy
from module_energy.controller.cost_controller import cost_month_summary, cost_month_top
from module_energy.dao.overview_dao import OverviewDao
from module_energy.entity.do.alert_event_do import EAlertEvent
from module_energy.entity.do.cost_recompute_record_do import ECostRecomputeRecord
from module_energy.entity.do.cost_record_do import ECostRecord
from module_energy.entity.do.report_archive_do import EReportArchive
from module_energy.entity.do.suggestion_do import ESuggestion
from module_energy.entity.do.suggestion_flow_log_do import ESuggestionFlowLog
from module_energy.entity.do.suggestion_verification_do import ESuggestionVerification
from module_energy.entity.vo.cost_vo import CostRecomputeRequest, CostTariffCreateRequest
from module_energy.entity.vo.report_vo import ReportPreviewRequest
from module_energy.entity.vo.suggestion_vo import ManualSuggestionCreateRequest
from module_energy.service.aggregation_service import AggregationService
from module_energy.service.baseline_service import BaselineService
from module_energy.service.cost_query_service import CostQueryService
from module_energy.service.cost_service import CostService
from module_energy.service.report_service import ReportService
from module_energy.service.rule_engine_service import RuleEngineService
from module_energy.service.suggestion_service import SuggestionService

_ROOT = Path(__file__).resolve().parents[2]
_DB_URL = 'mysql+asyncmy://demo:bdemo_dev@127.0.0.1:3306/{database}'
_MONEY = Decimal('0.01')
_V2 = 2
_AREA_B_ID = 2


class Act5ModuleInventoryContractTest(unittest.TestCase):
    def test_module_inventory_marks_cost_and_reports_implemented(self) -> None:
        self.assertIn('5.9 成本核算与报表    REQ-051~062（已实现）', module_energy.__doc__ or '')

    def test_workflow_calls_both_legacy_cost_controller_consumers(self) -> None:
        source = inspect.getsource(
            Act5WorkflowIntegrationTest.test_real_workflow_and_deterministic_reset
        )
        self.assertIn('cost_month_summary(', source)
        self.assertIn('cost_month_top(', source)


class Act5WorkflowIntegrationTest(unittest.IsolatedAsyncioTestCase):
    """One real v1 -> v2 -> report -> suggestion -> reset journey."""

    async def test_real_workflow_and_deterministic_reset(self) -> None:  # noqa: PLR0915
        database = self._required_database()
        restored = False
        engine = None
        db = None
        try:
            await self._reset_and_bootstrap(database)
            engine, db = self._session(database)
            fresh = await self._fresh_snapshot(db)
            self.assertEqual(
                {
                    'costs': 75,
                    'currentCosts': 75,
                    'recomputes': 0,
                    'archives': 0,
                    'suggestions': 8,
                    'flows': 25,
                    'verifications': 3,
                    'alerts': 20,
                    'r10': 1,
                },
                fresh['counts'],
            )
            self.assertEqual(
                {
                    '2026-05': Decimal('138014.50'),
                    '2026-06': Decimal('161920.03'),
                    '2026-07': Decimal('59501.99'),
                },
                fresh['electricity'],
            )
            self.assertEqual(Decimal('1313.12'), fresh['waterJuly'])
            self.assertEqual(Decimal('4030.33'), fresh['airJuly'])

            monthly_request = ReportPreviewRequest.model_validate(
                {
                    'templateCode': 'ENERGY_MONTHLY',
                    'period': '2026-06',
                    'filters': {'zone': 'ALL', 'energyType': 'electricity'},
                }
            )
            fresh_preview = await ReportService.preview(db, monthly_request)
            fresh_signature = fresh_preview['signature']
            v1_archive = await ReportService.create_archive(
                db,
                monthly_request,
                archived_by='finance_user',
            )
            self.assertEqual(fresh_signature, v1_archive['fullSignature'])
            v1_cost_items = v1_archive['payloadSnapshot']['sections']['costSection'][
                'items'
            ]
            self.assertTrue(v1_cost_items)
            self.assertEqual(
                {'v1'},
                {
                    item['costVersion']
                    for item in v1_archive['versionSnapshots']['costVersions']
                },
            )

            v1_rows = await self._batch_rows(db, version=1)
            v1_evidence = {
                int(row.id): (Decimal(row.total_cost), row.signature) for row in v1_rows
            }
            overview_before = await OverviewDao.get_cost_month_summary(
                db, '2026-06', 'electricity'
            )
            top_before = await OverviewDao.get_top_cost_objects(db, '2026-06', limit=5)

            tariff = await CostService.create_tariff_version(
                db,
                CostTariffCreateRequest.model_validate(
                    {
                        'energyType': 'electricity',
                        'effectiveFrom': '2026-06-01',
                        'prices': {'peak': '1.35', 'flat': '0.85', 'valley': '0.45'},
                        'remark': '第五幕端到端回归：六月电价版本生效',
                    }
                ),
                operator='finance_user',
            )
            self.assertEqual(2, tariff['tariffVersion'])
            recompute = await CostService.recompute(
                db,
                CostRecomputeRequest.model_validate(
                    {
                        'statMonth': '2026-06',
                        'energyType': 'electricity',
                        'triggerReason': '六月电价版本 v2 生效',
                        'tariffVersion': 2,
                        'allocRuleVersion': 1,
                    }
                ),
                operator='energy_mgr',
            )
            self.assertEqual((1, 2, 'pending'), (
                recompute.old_cost_version,
                recompute.new_cost_version,
                recompute.review_status,
            ))
            self.assertTrue(recompute.diff_summary)
            self.assertTrue(
                {'metric', 'oldValue', 'newValue', 'deltaPct'}
                <= set(recompute.diff_summary[0])
            )

            old_after = await self._batch_rows(db, version=1)
            self.assertEqual(
                v1_evidence,
                {
                    int(row.id): (Decimal(row.total_cost), row.signature)
                    for row in old_after
                },
            )
            self.assertTrue(all(not row.is_current for row in old_after))
            v2_rows = await self._batch_rows(db, version=2)
            self.assertEqual(len(v1_rows), len(v2_rows))
            self.assertTrue(all(row.is_current for row in v2_rows))
            self.assertEqual([], await self._invalid_current_groups(db))

            overview_after = await OverviewDao.get_cost_month_summary(
                db, '2026-06', 'electricity'
            )
            top_after = await OverviewDao.get_top_cost_objects(db, '2026-06', limit=5)
            expected_area_total = await self._current_area_total(db)
            self.assertEqual(expected_area_total, Decimal(overview_after['total_cost']))
            self.assertEqual(2, int(overview_after['current_cost_version']))
            self.assertNotEqual(
                Decimal(overview_before['total_cost']),
                Decimal(overview_after['total_cost']),
            )
            self.assertTrue(
                all(int(row['current_cost_version']) == _V2 for row in top_after)
            )
            self.assertNotEqual(
                [(row['equipment_id'], row['total_cost']) for row in top_before],
                [(row['equipment_id'], row['total_cost']) for row in top_after],
            )

            legacy_summary_response = await cost_month_summary(
                request=None,
                query_db=db,
                energy_type='electricity',
                stat_month='2026-06',
                area_id=None,
            )
            legacy_summary = json.loads(legacy_summary_response.body)['data']
            legacy_total = Decimal(str(legacy_summary['summary']['total_cost']))
            self.assertEqual(expected_area_total, legacy_total)
            self.assertEqual(await self._current_system_total(db), legacy_total)
            self.assertEqual('v2', legacy_summary['summary']['currentCostVersion'])

            legacy_top_response = await cost_month_top(
                request=None,
                query_db=db,
                stat_month='2026-06',
                limit=5,
                area_id=None,
            )
            legacy_top = json.loads(legacy_top_response.body)['data']['items']
            self.assertTrue(
                all(item['currentCostVersion'] == 'v2' for item in legacy_top)
            )
            self.assertEqual(
                [
                    (
                        int(row['equipment_id']),
                        Decimal(row['total_cost']).quantize(_MONEY),
                    )
                    for row in top_after
                ],
                [
                    (
                        int(item['equipment_id']),
                        Decimal(str(item['total_cost'])).quantize(_MONEY),
                    )
                    for item in legacy_top
                ],
            )

            rule_counts = await RuleEngineService.run_all(db)
            self.assertEqual(1, rule_counts['R10'])
            self.assertEqual((20, 1), await self._alert_counts(db))

            review = await CostService.review(
                db,
                recompute.recompute_id,
                'approve',
                '财务复核冻结证据与重算差异一致',
                reviewer='finance_user',
            )
            self.assertEqual('approved', review.review_status)

            month_view = await CostQueryService.get_month_view(
                db,
                stat_month='2026-06',
                zone='B',
                energy_type='electricity',
                group_by='area',
                focus='R10',
            )
            # INJ-06 的产品锚点是“约 12pp 抬升”；展示仍保留引擎冻结实值。
            self.assertEqual(12, round(month_view['anomalyEvidence']['diffPp']))
            self.assertEqual('v2', month_view['summary']['currentCostVersion'])
            area_b = next(
                row
                for row in v2_rows
                if row.object_type == 'area' and row.object_id == _AREA_B_ID
            )
            trace = await CostQueryService.get_trace(
                db,
                stat_month='2026-06',
                object_type='area',
                object_id=2,
                energy_type='electricity',
                cost_version=2,
            )
            self.assertEqual('notApplied', trace['allocationEvidence']['allocationStatus'])
            self.assertEqual([], trace['allocationEvidence']['allocationDetails'])
            self.assertEqual('本对象无共享分摊', trace['allocationEvidence']['message'])
            self.assertEqual(recompute.recompute_id, trace['recomputeChain'][0]['recomputeId'])

            current_preview = await ReportService.preview(db, monthly_request)
            self.assertNotEqual(fresh_signature, current_preview['signature'])
            self.assertNotEqual(
                v1_cost_items,
                current_preview['sections']['costSection']['items'],
            )
            exported = await ReportService.export(db, monthly_request)
            self.assertEqual(current_preview['signature'], exported['signature'])
            workbook = load_workbook(BytesIO(exported['content']))
            self.assertIn('口径说明', workbook.sheetnames)
            self.assertTrue(
                all(workbook[name].freeze_panes == 'A7' for name in workbook.sheetnames if name != '口径说明')
            )
            current_archive = await ReportService.create_archive(
                db,
                monthly_request,
                archived_by='finance_user',
            )
            self.assertEqual(current_preview['signature'], current_archive['fullSignature'])
            frozen_v1 = await ReportService.get_archive(db, v1_archive['archiveId'])
            frozen_v1_file = await ReportService.export_archive(db, v1_archive['archiveId'])
            self.assertEqual(fresh_signature, frozen_v1['fullSignature'])
            self.assertEqual(fresh_signature, frozen_v1_file['signature'])
            self.assertEqual(
                v1_archive['payloadSnapshot'],
                frozen_v1['payloadSnapshot'],
            )
            self.assertEqual(
                v1_archive['versionSnapshots'],
                frozen_v1['versionSnapshots'],
            )
            self.assertEqual(
                v1_cost_items,
                frozen_v1['payloadSnapshot']['sections']['costSection']['items'],
            )
            frozen_workbook = load_workbook(BytesIO(frozen_v1_file['content']))
            frozen_cost_sheet = frozen_workbook['costSection']
            frozen_headers = {
                frozen_cost_sheet.cell(6, column).value: column
                for column in range(1, frozen_cost_sheet.max_column + 1)
            }
            self.assertEqual(
                float(v1_cost_items[0]['totalCost']),
                frozen_cost_sheet.cell(7, frozen_headers['totalCost']).value,
            )
            self.assertEqual(
                v1_cost_items[0]['currentCostVersion'],
                frozen_cost_sheet.cell(
                    7,
                    frozen_headers['currentCostVersion'],
                ).value,
            )

            r10_event_id = int(
                (
                    await db.execute(
                        select(EAlertEvent.event_id)
                        .where(EAlertEvent.rule_code == 'R10')
                        .limit(1)
                    )
                ).scalar_one()
            )
            suggestion_request = ManualSuggestionCreateRequest.model_validate(
                {
                    'title': '复核 B 区峰段成本异常',
                    'sourceDescription': '财务月度成本复核发现峰段构成异常',
                    'objectType': 'area',
                    'objectId': 2,
                    'areaId': 2,
                    'measureContent': '人工核对峰段用量、单价版本与分摊口径',
                    'templateId': 1,
                    'priorityFactors': {
                        'energyScale': 80,
                        'costImpact': 75,
                        'duration': 60,
                        'implementationDifficulty': 30,
                        'safetyImpact': 20,
                    },
                    'sourceContext': {
                        **trace['suggestionContext'],
                        'recomputeId': recompute.recompute_id,
                        'alertEventId': r10_event_id,
                        'costSignature': area_b.signature,
                    },
                }
            )
            suggestion = await SuggestionService.create_manual(
                db,
                request=suggestion_request,
                operator='finance_user',
                operator_role='finance',
                role_keys={'finance'},
            )
            self.assertTrue(suggestion['created'])
            self.assertEqual(
                suggestion['suggestion']['suggestionId'],
                suggestion['relatedSuggestionId'],
            )
            trace_after = await CostQueryService.get_trace(
                db,
                stat_month='2026-06',
                object_type='area',
                object_id=2,
                energy_type='electricity',
                cost_version=2,
            )
            self.assertEqual(suggestion['relatedSuggestionId'], trace_after['relatedSuggestionId'])
            self.assertEqual(
                {'costs': 90, 'v1': 75, 'v2': 15, 'diffs': 1, 'archives': 2, 'suggestions': 9},
                await self._dirty_counts(db),
            )

            await db.close()
            await engine.dispose()
            db = None
            engine = None
            await self._reset_and_bootstrap(database)
            engine, db = self._session(database)
            restored_snapshot = await self._fresh_snapshot(db)
            restored_preview = await ReportService.preview(db, monthly_request)
            self.assertEqual(fresh['counts'], restored_snapshot['counts'])
            self.assertEqual(fresh['electricity'], restored_snapshot['electricity'])
            self.assertEqual(fresh_signature, restored_preview['signature'])
            self.assertEqual([], await self._invalid_current_groups(db))
            restored = True
        finally:
            if db is not None:
                await db.close()
            if engine is not None:
                await engine.dispose()
            if not restored:
                await self._reset_and_bootstrap(database)

    @staticmethod
    def _required_database() -> str:
        act4 = os.environ.get('ACT4_TEST_DB', '')
        act5 = os.environ.get('ACT5_TEST_DB', '')
        if not act4 or not act5:
            raise AssertionError('ACT4_TEST_DB and ACT5_TEST_DB are both required; skip is forbidden')
        if act4 != act5:
            raise AssertionError('ACT4_TEST_DB and ACT5_TEST_DB must point to the same database')
        if not act5.startswith('codex_') or not act5.endswith('_test'):
            raise AssertionError('workflow database must match codex_*_test')
        return act5

    @staticmethod
    def _session(database: str) -> tuple[Any, Any]:
        engine = create_async_engine(_DB_URL.format(database=database))
        return engine, async_sessionmaker(engine, expire_on_commit=False)()

    @classmethod
    async def _reset_and_bootstrap(cls, database: str) -> None:
        env = {**os.environ, 'ACT4_TEST_DB': database, 'ACT5_TEST_DB': database, 'PYTHONPATH': 'backend'}

        def run_reset() -> None:
            subprocess.run(
                [str(_ROOT / 'backend/.venv/bin/python'), str(_ROOT / 'datagen/generate_demo_data.py'), '--reset'],
                cwd=_ROOT,
                env=env,
                check=True,
                capture_output=True,
                text=True,
            )

        await asyncio.to_thread(run_reset)
        engine, db = cls._session(database)
        try:
            await AggregationService.rebuild_all(db)
            await BaselineService.compute_and_publish(db, force_republish=True)
            await BaselineService.compute_daily_deviation(db)
            await CostService.initialize_v1(db, computed_by='pipeline-bootstrap')
            await CostService.validate_current_integrity(db)
            counts = await RuleEngineService.run_all(db)
            if counts.get('R10') != 1:
                raise AssertionError(f'bootstrap R10 count drifted: {counts!r}')
        finally:
            await db.close()
            await engine.dispose()

    @staticmethod
    async def _batch_rows(db: Any, *, version: int) -> list[ECostRecord]:
        return list(
            (
                await db.execute(
                    select(ECostRecord)
                    .where(
                        ECostRecord.stat_month == '2026-06',
                        ECostRecord.energy_type_code == 'electricity',
                        ECostRecord.cost_version == version,
                    )
                    .order_by(ECostRecord.id)
                )
            ).scalars()
        )

    @staticmethod
    async def _invalid_current_groups(db: Any) -> list[tuple[Any, ...]]:
        return list(
            (
                await db.execute(
                    text(
                        """
                        SELECT object_type, IFNULL(object_id, 0), stat_month,
                               energy_type_code, SUM(is_current) AS current_count
                        FROM e_cost_record
                        GROUP BY object_type, IFNULL(object_id, 0), stat_month,
                                 energy_type_code
                        HAVING current_count <> 1
                        """
                    )
                )
            ).all()
        )

    @staticmethod
    async def _current_area_total(db: Any) -> Decimal:
        value = (
            await db.execute(
                select(func.sum(ECostRecord.total_cost)).where(
                    ECostRecord.stat_month == '2026-06',
                    ECostRecord.energy_type_code == 'electricity',
                    ECostRecord.object_type == 'area',
                    ECostRecord.is_current.is_(True),
                    ECostRecord.status != 'void',
                )
            )
        ).scalar_one()
        return Decimal(value)

    @staticmethod
    async def _current_system_total(db: Any) -> Decimal:
        value = (
            await db.execute(
                select(ECostRecord.total_cost).where(
                    ECostRecord.stat_month == '2026-06',
                    ECostRecord.energy_type_code == 'electricity',
                    ECostRecord.object_type == 'system',
                    ECostRecord.object_id == 0,
                    ECostRecord.is_current.is_(True),
                    ECostRecord.status != 'void',
                )
            )
        ).scalar_one()
        return Decimal(value)

    @staticmethod
    async def _alert_counts(db: Any) -> tuple[int, int]:
        total = int((await db.execute(select(func.count()).select_from(EAlertEvent))).scalar_one())
        r10 = int(
            (
                await db.execute(
                    select(func.count()).select_from(EAlertEvent).where(EAlertEvent.rule_code == 'R10')
                )
            ).scalar_one()
        )
        return total, r10

    @classmethod
    async def _fresh_snapshot(cls, db: Any) -> dict[str, Any]:
        models = {
            'costs': ECostRecord,
            'recomputes': ECostRecomputeRecord,
            'archives': EReportArchive,
            'suggestions': ESuggestion,
            'flows': ESuggestionFlowLog,
            'verifications': ESuggestionVerification,
            'alerts': EAlertEvent,
        }
        counts = {
            name: int((await db.execute(select(func.count()).select_from(model))).scalar_one())
            for name, model in models.items()
        }
        counts['currentCosts'] = int(
            (
                await db.execute(
                    select(func.count()).select_from(ECostRecord).where(ECostRecord.is_current.is_(True))
                )
            ).scalar_one()
        )
        counts['r10'] = int(
            (
                await db.execute(
                    select(func.count()).select_from(EAlertEvent).where(EAlertEvent.rule_code == 'R10')
                )
            ).scalar_one()
        )
        electricity_rows = (
            await db.execute(
                select(ECostRecord.stat_month, ECostRecord.total_cost).where(
                    ECostRecord.object_type == 'system',
                    ECostRecord.object_id == 0,
                    ECostRecord.energy_type_code == 'electricity',
                    ECostRecord.is_current.is_(True),
                )
            )
        ).all()
        media = {}
        for energy in ('water', 'compressed_air'):
            media[energy] = Decimal(
                (
                    await db.execute(
                        select(ECostRecord.total_cost).where(
                            ECostRecord.object_type == 'system',
                            ECostRecord.object_id == 0,
                            ECostRecord.stat_month == '2026-07',
                            ECostRecord.energy_type_code == energy,
                            ECostRecord.is_current.is_(True),
                        )
                    )
                ).scalar_one()
            ).quantize(_MONEY)
        return {
            'counts': counts,
            'electricity': {
                month: Decimal(cost).quantize(_MONEY) for month, cost in electricity_rows
            },
            'waterJuly': media['water'],
            'airJuly': media['compressed_air'],
        }

    @staticmethod
    async def _dirty_counts(db: Any) -> dict[str, int]:
        total = int((await db.execute(select(func.count()).select_from(ECostRecord))).scalar_one())
        versions = dict(
            (
                await db.execute(
                    select(ECostRecord.cost_version, func.count())
                    .group_by(ECostRecord.cost_version)
                    .order_by(ECostRecord.cost_version)
                )
            ).all()
        )
        return {
            'costs': total,
            'v1': int(versions.get(1, 0)),
            'v2': int(versions.get(2, 0)),
            'diffs': int((await db.execute(select(func.count()).select_from(ECostRecomputeRecord))).scalar_one()),
            'archives': int((await db.execute(select(func.count()).select_from(EReportArchive))).scalar_one()),
            'suggestions': int((await db.execute(select(func.count()).select_from(ESuggestion))).scalar_one()),
        }
